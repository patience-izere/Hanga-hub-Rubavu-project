from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Hanga-hub-Rubavu-project")
OUTPUT = ROOT / "output" / "docx" / "OPedu_AR_Platform_Evaluation.docx"

NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5B6573"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_GREEN = "E7F2EC"
PALE_GOLD = "FFF4D6"
PALE_RED = "FCE8E6"
WHITE = "FFFFFF"
BLACK = "1A1A1A"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_table_borders(table, color="C7CDD4", size="5"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[idx] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def set_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_field(paragraph, field_code):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_font(run, size=9, color=MUTED)


def add_list_numbering(document):
    numbering = document.part.numbering_part.element
    existing_abstract = [int(e.get(qn("w:abstractNumId"))) for e in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(e.get(qn("w:numId"))) for e in numbering.findall(qn("w:num"))]

    def create(ordered):
        abstract_id = max(existing_abstract + [0]) + 1
        existing_abstract.append(abstract_id)
        num_id = max(existing_num + [0]) + 1
        existing_num.append(num_id)
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "decimal" if ordered else "bullet")
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), "%1." if ordered else "\uf0b7")
        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "left")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "720")
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "360")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "160")
        spacing.set(qn("w:line"), "280")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.extend([tabs, ind, spacing])
        level.extend([start, num_fmt, lvl_text, lvl_jc, p_pr])
        if not ordered:
            r_pr = OxmlElement("w:rPr")
            r_fonts = OxmlElement("w:rFonts")
            r_fonts.set(qn("w:ascii"), "Symbol")
            r_fonts.set(qn("w:hAnsi"), "Symbol")
            r_pr.append(r_fonts)
            level.append(r_pr)
        abstract.append(level)
        first_num = numbering.find(qn("w:num"))
        if first_num is None:
            numbering.append(abstract)
        else:
            numbering.insert(numbering.index(first_num), abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)
        return num_id

    return create(False), create(True)


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])
    p_pr.append(num_pr)


doc = Document()
doc.settings.odd_and_even_pages_header_footer = False
section = doc.sections[0]
section.different_first_page_header_footer = False
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1

for name, size, color, before, after in (
    ("Title", 26, NAVY, 0, 6),
    ("Subtitle", 13, MUTED, 0, 12),
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, DARK_BLUE, 8, 4),
):
    style = styles[name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = name != "Subtitle"
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

title_style_p_pr = styles["Title"]._element.get_or_add_pPr()
title_style_border = title_style_p_pr.find(qn("w:pBdr"))
if title_style_border is not None:
    title_style_p_pr.remove(title_style_border)

bullet_num_id, decimal_num_id = add_list_numbering(doc)


def add_heading(text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    return p


def add_para(text="", bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        first = p.add_run(bold_lead)
        set_font(first, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_font(rest)
    else:
        run = p.add_run(text)
        set_font(run)
    return p


def add_bullet(text):
    p = doc.add_paragraph()
    apply_num(p, bullet_num_id)
    set_font(p.add_run(text))
    return p


def add_number(text, num_id=None):
    p = doc.add_paragraph()
    apply_num(p, num_id or decimal_num_id)
    set_font(p.add_run(text))
    return p


def add_callout(label, text, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table, color=fill, size="4")
    set_repeat_table_header(table.rows[0])
    set_row_cant_split(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    lead = p.add_run(f"{label}: ")
    set_font(lead, bold=True, color=NAVY)
    set_font(p.add_run(text), color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_table(headers, rows, widths, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, NAVY)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_font(p.add_run(header), size=font_size, color=WHITE, bold=True)
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        set_row_cant_split(table.rows[-1])
        if row_idx % 2 == 1:
            for cell in cells:
                set_cell_shading(cell, "FAFBFC")
        for idx, value in enumerate(row):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            set_font(p.add_run(str(value)), size=font_size)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


# Running header and footer
header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hp.paragraph_format.space_after = Pt(0)
set_font(hp.add_run("OPedu | AR Platform Evaluation"), size=9, color=MUTED, bold=True)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fp.paragraph_format.space_before = Pt(0)
set_font(fp.add_run("Evaluation report  |  Page "), size=9, color=MUTED)
add_field(fp, "PAGE")
set_font(fp.add_run(" of "), size=9, color=MUTED)
add_field(fp, "NUMPAGES")

# Cover
doc.add_paragraph().paragraph_format.space_after = Pt(34)
kicker = doc.add_paragraph()
kicker.paragraph_format.space_after = Pt(8)
set_font(kicker.add_run("TECHNICAL PRODUCT EVALUATION"), size=10, color=BLUE, bold=True)
title = doc.add_paragraph("OPedu AR Platform Evaluation", style="Title")
title_p_pr = title._p.get_or_add_pPr()
title_border = title_p_pr.find(qn("w:pBdr"))
if title_border is not None:
    title_p_pr.remove(title_border)
subtitle = doc.add_paragraph(
    "Assessment against the low-bandwidth WebAR e-learning research framework in "
    "25093891 AR ELearning Dissertation",
    style="Subtitle",
)
subtitle.paragraph_format.space_after = Pt(22)

meta = add_table(
    ["Evaluation item", "Detail"],
    [
        ("Platform", "OPedu technical learning platform"),
        ("Evidence date", "22 August 2026"),
        ("Scope", "Current repository, dissertation requirements, AR readiness, and pilot pathway"),
        ("Decision requested", "Prioritize the work that converts the desktop 3D learning system into a defensible WebAR platform"),
    ],
    [2200, 7160],
    font_size=10,
)
doc.add_paragraph().paragraph_format.space_after = Pt(12)
add_callout(
    "Headline verdict",
    "OPedu is already a credible secure learning and 3D assessment foundation. It is not yet an augmented-reality platform. The shortest responsible path is to preserve the existing curriculum, evidence, and grading core while adding a shared WebAR renderer, capability-aware fallback, low-bandwidth delivery, xAPI-compatible events, rules-based adaptation, and a measured school pilot.",
    PALE_GOLD,
)

doc.add_page_break()

add_heading("Executive evaluation", 1)
add_para(
    "The dissertation proposes a browser-delivered, low-bandwidth AR learning system that combines "
    "WebAR, adaptive content sequencing, standardized learning analytics, accessible fallback, "
    "content authoring, and empirical evaluation. OPedu now implements much of the non-AR platform "
    "beneath that concept: React delivery, Django APIs, PostgreSQL learning records, institutional roles, "
    "versioned course and simulation content, a desktop 3D battery procedure, deterministic scoring, "
    "immutable evidence, competency results, and instructor review."
)
add_para(
    "The central gap is categorical, not cosmetic: the current Three.js scene is a virtual workshop on a "
    "screen. It does not overlay content on a live camera view, recognize markers or surfaces, create "
    "spatial anchors, or enter a WebXR immersive-AR session. Calling the current product an AR platform "
    "would therefore be premature."
)
add_callout(
    "Recommended product position now",
    "Describe OPedu as a competency-based technical learning platform with a browser 3D simulation prototype and an AR-ready data model. Reserve 'WebAR platform' for the release that passes the AR qualification gate in Section 6.",
    LIGHT_BLUE,
)

add_heading("Indicative maturity", 2)
add_para(
    "Scores below are an evaluation judgment, not a statistical measure: 0 = absent, 1 = planned, "
    "2 = working foundation, 3 = pilot-capable, and 4 = validated in real use."
)
add_table(
    ["Capability", "Score", "Evaluation"],
    [
        ("Secure learning-platform foundation", "3 / 4", "Working roles, school isolation, curriculum, assignments, attempts, evidence, and administration."),
        ("Desktop 3D technical simulation", "2 / 4", "A guided battery scene works, but assets and interactions are still prototype-grade."),
        ("True WebAR delivery", "0 / 4", "No camera overlay, spatial tracking, anchors, hit testing, or immersive-AR session."),
        ("Low-bandwidth and offline operation", "1 / 4", "Browser delivery and route splitting help; caching, offline evidence, optimized assets, and network gates remain."),
        ("Adaptive sequencing", "1 / 4", "Prerequisites and deterministic next steps exist; learner-specific recommendation logic does not."),
        ("xAPI/LRS analytics", "1 / 4", "Ordered immutable events exist, but stable xAPI identifiers, statements, export, and LRS integration do not."),
        ("Research and pilot readiness", "1 / 4", "A roadmap exists; instruments, consent workflow, device study, field pilot, and actual outcome data remain."),
    ],
    [2500, 950, 5910],
    font_size=8.5,
)

add_heading("1. Scope, method, and evidence quality", 1)
add_para(
    "This report treats the attached dissertation as research evidence and the repository as implementation "
    "evidence. Instructions found inside the dissertation were not treated as user instructions. No product "
    "code was changed for this evaluation."
)
add_heading("Evidence reviewed", 2)
add_bullet("The complete 48-page dissertation, including its literature review, conceptual framework, methodology, architecture, proposed implementation stack, planned evaluation, limitations, recommendations, and appendices.")
add_bullet("The current OPedu README, architecture decisions, assessment/content-versioning documentation, pilot notes, and execution backlog.")
add_bullet("The current frontend package and core learner simulation implementation, including the React Three Fiber battery workshop and accessible HTML action controls.")
add_bullet("The current Django learning-domain models and API surfaces for curriculum, content governance, attempts, events, grading, step results, competency results, auditing, and school roles.")

add_heading("Important research caveat", 2)
add_para(
    "The dissertation is explicit that its pilot had not been conducted at the time of writing. Its Chapters "
    "Four and Five provide a prototype specification and planned reporting structure, not measured evidence "
    "that the proposed system is usable, accepted, accurate, or educationally effective in Rwanda. Its value "
    "is therefore strongest as a requirements and evaluation framework. Product claims must wait for actual "
    "field results."
)

add_heading("2. The target described by the dissertation", 1)
add_para(
    "The study's target is broader than a camera-enabled 3D viewer. It is a five-layer learning system in "
    "which AR delivery, authoring, assessment, adaptive sequencing, analytics, institutional integration, and "
    "fallback modes share one data model."
)
add_table(
    ["Research requirement", "Meaning for OPedu"],
    [
        ("Browser-based WebAR", "Run through a standard mobile browser and overlay digital guidance on a real object or workspace without a native app install."),
        ("Low-bandwidth delivery", "Use small, cacheable assets, quality tiers, resilient retries, offline progress, and tests under realistic 3G/4G conditions."),
        ("Adaptive sequencing", "Choose the next lesson, remediation, or challenge from mastery gap, engagement, cognitive-load proxy, and device reliability while respecting prerequisites."),
        ("xAPI learning analytics", "Represent learning as stable actor-verb-object-result statements, with deduplication and optional export to a compliant Learning Record Store."),
        ("2D/video fallback", "Give learners on unsupported devices an equivalent learning and assessment path, not a reduced afterthought."),
        ("Authoring", "Enable instructors or content authors to create and publish AR-enriched modules without editing source code."),
        ("Evaluation", "Measure recognition accuracy, usability, acceptance, learning gain, reliability, device fairness, and longer-term transfer."),
    ],
    [2500, 6860],
    font_size=9.5,
)

add_heading("Research-derived success conditions", 2)
add_bullet("Perceived usefulness, ease of use, and enjoyment must be designed and measured; novelty alone is not success.")
add_bullet("Cognitive load must remain controlled through simple interfaces, progressive prompts, and recovery from recognition failure.")
add_bullet("The system must work on representative mid-range Android devices and must not disadvantage users without WebXR support.")
add_bullet("Recognition and recommendation accuracy must be measured separately from learner performance.")
add_bullet("Real classroom evidence is required because the literature is dominated by small, short, single-institution studies.")

add_heading("3. What OPedu has today", 1)
add_heading("A. Platform and governance foundation - strong", 2)
add_bullet("React 19, TypeScript, Vite, React Router, TanStack Query, Three.js, React Three Fiber, and Drei provide a modern browser client and a shared 3D foundation.")
add_bullet("Django 5.2 and Django REST Framework provide authenticated APIs and administration; PostgreSQL is the intended system of record and Redis is available for future asynchronous work.")
add_bullet("School membership and role controls support administrators, content authors, instructors, and learners while preserving school isolation.")
add_bullet("Programs, courses, modules, competencies, lessons, cohorts, enrollment, prerequisites, publication states, and content versions establish a more complete education domain than the dissertation's simplified ERD.")

add_heading("B. Technical-learning and assessment foundation - strong", 2)
add_bullet("The first vertical lesson is an automotive battery inspection and diagnosis procedure with ordered steps, safety-critical actions, tools, hazards, hints, tolerances, acceptable actions, and feedback rules.")
add_bullet("Published scenario, grading-policy, and asset-package versions are immutable, which preserves reproducibility after authoring changes.")
add_bullet("Attempt events are ordered and immutable. Completion creates immutable step and competency results with evidence references, safety violations, hint use, duration, tolerance outcomes, mastery thresholds, and final outcomes.")
add_bullet("An instructor can review completion metrics, safety flags, scores, and the event timeline for learners in the same school.")

add_heading("C. 3D interaction foundation - working prototype", 2)
add_bullet("The learner can orbit and zoom a Three.js workshop, select battery-related objects, receive immediate guidance, resume progress, and submit the attempt for server-side scoring.")
add_bullet("Equivalent HTML buttons expose the same action vocabulary as the 3D objects, creating a valuable accessibility and fallback pattern that should be retained in AR.")
add_bullet("The 3D route is lazy-loaded, reducing the cost imposed on public and dashboard pages.")
add_para(
    "However, the scene uses generated geometric primitives rather than a licensed, optimized, "
    "versioned real equipment model. It is a useful interaction demonstrator, not yet production technical content."
)

add_heading("4. Research-to-repository traceability", 1)
add_table(
    ["Dissertation capability", "Current status", "Repository evidence", "Gap to target"],
    [
        ("Browser client", "Implemented", "React/Vite application; all public routes are React-owned.", "Mobile and headset-browser validation remains."),
        ("3D rendering", "Prototype", "React Three Fiber battery workshop with orbit/zoom and selectable objects.", "Production glTF, loader, interaction framework, budgets, and physical-device testing."),
        ("True WebAR", "Missing", "No WebXR AR session or camera/spatial APIs found.", "Camera passthrough, hit test, anchors, marker/surface recognition, placement, session lifecycle."),
        ("Authentication/users", "Implemented", "Session authentication, invitations, roles, lockout, school isolation.", "Pilot consent and research participant state are separate requirements."),
        ("Course/content model", "Implemented", "Programs, modules, competencies, lessons, scenarios, tools, hazards, hints, tolerances.", "Low/no-code AR authoring and asset-validation UI."),
        ("Assessment/progress", "Implemented", "Assignments, attempts, immutable events, deterministic grading, step/competency results.", "Retry/history/remediation and instructor-authored observations."),
        ("Adaptive sequencing", "Foundation only", "Prerequisites, current-step guidance, mastery data.", "A transparent rules engine using mastery gaps and device/engagement signals."),
        ("Learning analytics", "Foundation only", "Ordered event evidence and instructor overview.", "Stable activity IDs, xAPI vocabulary, statement IDs, deduplication, export/LRS option."),
        ("2D fallback", "Partial", "Accessible HTML actions mirror the 3D action vocabulary.", "Capability-aware 2D/video lesson renderer with equivalent outcomes and parity tests."),
        ("Low-bandwidth", "Early", "Web delivery and lazy route loading.", "PWA, IndexedDB event queue, download/cache, optimized assets, quality tiers, 3G/4G budgets."),
        ("Evaluation", "Planned", "Pilot backlog and deterministic evidence records.", "TAM/SUS, pre/post tests, recognition metrics, consent, device study, analysis pipeline."),
        ("LMS/LRS integration", "Missing", "No SCORM, LTI, xAPI export, or LRS connector found.", "Choose the minimum institutional integration after pilot needs are confirmed."),
    ],
    [1900, 1100, 3150, 3210],
    font_size=8.2,
)

add_heading("5. What is missing to reach an AR platform", 1)
add_heading("5.1 AR runtime and physical-world registration", 2)
add_bullet("WebXR capability detection and an explicit, permission-aware Enter AR / Exit AR lifecycle.")
add_bullet("Camera passthrough with a secure-origin and browser-support strategy.")
add_bullet("Surface hit testing and spatial anchors for placing instructional content consistently in a learner's physical workspace.")
add_bullet("A marker-based fallback for devices or browsers that cannot provide reliable markerless tracking.")
add_bullet("Tracking-loss, poor-lighting, unsupported-device, camera-denied, and unsafe-workspace recovery flows.")
add_bullet("AR interaction telemetry: recognition success/failure, latency, re-scan count, tracking loss, placement confidence, and fallback use.")

add_heading("5.2 Shared assessment semantics", 2)
add_para(
    "Desktop, accessible HTML, marker AR, and markerless AR should invoke the same scenario action codes and "
    "the same server-side grading rules. Mode and recognition quality should be contextual evidence, never a "
    "separate grading engine. A learner must not lose marks because the browser failed to track a surface."
)

add_heading("5.3 Production technical content", 2)
add_bullet("One licensed and instructor-validated battery or electrical training-rig asset package, with real dimensions, named parts, pivots, interaction points, collision/selection volumes, and safety zones.")
add_bullet("glTF validation, Meshopt or Draco geometry compression where supported, KTX2 textures, checksums, manifests, thumbnails, and explicit asset budgets.")
add_bullet("Content validation by qualified technical instructors against an approved curriculum standard, including hazards, sequence, tolerances, remediation, and physical transfer tasks.")
add_bullet("A safe AR lesson design: begin with mock-ups or de-energized training rigs; do not position the product as a substitute for supervised workshop practice.")

add_heading("5.4 Low-bandwidth and device resilience", 2)
add_bullet("Installable PWA shell, per-lesson download, versioned cache, checksum verification, and clear storage controls.")
add_bullet("IndexedDB storage for attempts and events, client-generated event IDs, ordered idempotent synchronization, retry/backoff, and conflict preservation.")
add_bullet("High, medium, and low rendering tiers selected from measured device capability rather than device brand alone.")
add_bullet("Test budgets for first load, cached load, frame rate, memory, recognition latency, and evidence synchronization under institutional Wi-Fi and constrained 3G/4G profiles.")

add_heading("5.5 Adaptive learning and standardized analytics", 2)
add_bullet("Stable URI-based identifiers for actors, verbs, activities, competencies, scenarios, steps, and results.")
add_bullet("An internal xAPI statement adapter that maps immutable OPedu events to actor-verb-object-result-context statements and assigns globally unique statement IDs.")
add_bullet("Deduplication and export first; deploy a separate LRS only when an institution or research protocol requires it.")
add_bullet("A transparent rules-based recommender using mastery gaps, prerequisite completion, attempts, hints, safety errors, and device reliability. Machine learning should wait until sufficient consented pilot data exist.")

add_heading("5.6 Authoring, instructor operations, and evidence", 2)
add_bullet("A low/no-code content-authoring workflow for scenes, anchors, actions, prompts, tools, hazards, tolerances, fallback media, and competency mappings.")
add_bullet("Scenario preview without assessment logging; publishing checks that reject missing fallback, oversized assets, broken identifiers, and unsupported interactions.")
add_bullet("Instructor observations and published feedback, cohort assignment, review filters, competency aggregation, and CSV/research export.")
add_bullet("TAM/SUS questionnaires, pre/post knowledge checks, practical transfer rubrics, research-consent records, and pseudonymized analysis exports.")

add_heading("6. AR qualification gate", 1)
add_para("OPedu should call a release an AR platform only when all of the following are demonstrated:")
_, ar_gate_num_id = add_list_numbering(doc)
add_number("A supported mobile browser can enter an AR session and display the real camera environment.", ar_gate_num_id)
add_number("The learner can register a training object or surface and keep guidance spatially aligned during the task.", ar_gate_num_id)
add_number("At least one real technical-school procedure uses AR-specific overlays that add instructional value beyond the desktop 3D scene.", ar_gate_num_id)
add_number("The same procedure can be completed through a first-class 2D/desktop fallback with equivalent required evidence and grading.", ar_gate_num_id)
add_number("Tracking failure, low light, permission denial, and unsupported hardware have tested recovery paths.", ar_gate_num_id)
add_number("Assets and event synchronization meet approved performance budgets on the lowest supported device and constrained network profile.", ar_gate_num_id)
add_number("AR recognition quality is recorded separately from learner competency and is visible in evaluation exports.", ar_gate_num_id)
add_number("A qualified technical instructor has approved the content, hazards, tolerances, and real-workshop transfer task.", ar_gate_num_id)

add_heading("7. Recommended target architecture", 1)
add_para(
    "Retain the modular Django monolith and React client. The dissertation allows Django/Python, and OPedu's "
    "existing domain model is already stronger than replacing it with a new Node service. Add AR as a client "
    "capability and analytics as adapters around the existing evidence core, not as a rewrite."
)
add_table(
    ["Layer", "Recommended responsibility"],
    [
        ("Experience", "React application with a capability gateway that selects desktop 3D, accessible 2D, marker AR, or markerless WebXR AR."),
        ("Shared renderer", "React Three Fiber scene graph, one action vocabulary, interaction adapters by mode, progressive asset loading, and explicit AR session lifecycle."),
        ("Learning services", "Current Django APIs for identity, content, assignment, attempts, grading, results, instructor review, and authoring governance."),
        ("Adaptive service", "A deterministic policy module inside Django initially; version its inputs, weights, explanation, and recommendation result."),
        ("Evidence adapter", "Map OPedu events to xAPI-compatible statements; preserve PostgreSQL as the authoritative source and support optional LRS export."),
        ("Offline layer", "Service worker for shell/assets plus IndexedDB for content metadata, active attempts, and an idempotent event outbox."),
        ("Data/operations", "PostgreSQL, object storage/CDN for versioned assets, Redis/background jobs where needed, monitoring, and pseudonymized research exports."),
    ],
    [1900, 7460],
    font_size=9.4,
)

add_callout(
    "Architecture decision",
    "Do not introduce microservices, a new backend language, or a mandatory LRS merely to match the dissertation diagram. The research specifies logical services, not a requirement that every box be independently deployed.",
    PALE_GREEN,
)

add_heading("8. Technical-school AR use cases", 1)
add_table(
    ["Use case", "AR value", "Evidence to capture", "Safety boundary"],
    [
        ("Battery inspection", "Overlay terminal identity, meter setup, safe connection order, and voltage range on a training battery.", "Placement, action order, measurement, hints, safety errors, recognition quality.", "Use de-energized or instructor-approved training rigs; no unsupervised live-vehicle work."),
        ("Electrical wiring", "Identify components and show the next verified connection on a physical panel.", "Component recognition, connection sequence, continuity/voltage tolerance, correction path.", "Isolate power until instructor authorizes energization."),
        ("Mechanical assembly", "Align part orientation, fastener sequence, and torque specification with the real assembly.", "Part/anchor identity, order, torque range, rework, duration.", "Use guards, tool checks, and instructor-controlled machinery."),
        ("Plumbing/HVAC diagnosis", "Label flow direction, valves, gauges, and fault points on a training installation.", "Inspection choices, readings, diagnosis, remediation.", "No pressurized or hazardous system intervention without supervision."),
        ("Construction safety", "Reveal hazard zones, PPE requirements, and inspection points in a staged workspace.", "Hazard recognition, response time, missed hazards, feedback.", "AR supplements, never replaces, site induction and supervisor authority."),
    ],
    [1700, 2750, 2800, 2110],
    font_size=8.2,
)

add_heading("Recommended first AR module", 2)
add_para(
    "Convert the existing battery lesson rather than starting a new trade. Use a printed fiducial marker or "
    "known training rig for the first controlled version, then add markerless surface placement after the "
    "evidence and fallback path are stable. This reuses OPedu's published steps, tools, hazards, tolerances, "
    "grading policy, and instructor timeline while introducing only the AR-specific layer."
)

add_heading("9. Prioritized way forward", 1)
add_table(
    ["Phase", "Outcome", "Key deliverables", "Exit gate"],
    [
        ("0. Definition and validation", "A controlled, assessable AR brief", "Curriculum source; qualified instructor; battery training rig; supported device baseline; consent/data map; AR value hypothesis.", "Signed content, safety, device, and measurement brief."),
        ("1. AR minimum vertical slice", "One battery step in real WebAR", "Capability check; Enter/Exit AR; camera; marker or hit test; anchor; tracking recovery; same action API; 2D fallback.", "AR qualification items 1-5 pass on target Android devices."),
        ("2. Production content and delivery", "A complete low-bandwidth AR lesson", "Licensed glTF; compression; manifest; typed loader; quality tiers; asset cache; performance telemetry; full guided procedure.", "Load, frame-rate, memory, and recognition budgets pass."),
        ("3. Offline evidence and analytics", "Reliable school operation", "PWA; IndexedDB outbox; idempotent batch sync; stable xAPI IDs; statement export; recognition metrics; instructor visibility.", "An offline attempt synchronizes exactly once with complete evidence."),
        ("4. Adaptive and authoring workflow", "Maintainable multi-lesson platform", "Rules-based recommendations; explanation; authoring/preview; fallback validation; instructor observations and feedback.", "Authors publish a new validated scenario without source-code edits."),
        ("5. Pilot and evaluation", "Evidence for product and research claims", "Two sections or schools; 60-120 learners where feasible; TAM/SUS; pre/post; practical transfer; device/network analysis; interviews.", "Predefined usability, safety, parity, learning, and operational thresholds pass."),
        ("6. Scale decision", "Evidence-led expansion", "Retention study; multi-school replication; LMS/LRS integration if required; second trade; procurement and support model.", "Independent go/no-go review based on pilot evidence."),
    ],
    [1200, 1900, 3900, 2360],
    font_size=8.2,
)

add_heading("Priority order", 2)
_, priority_num_id = add_list_numbering(doc)
add_number("Validate one real technical procedure and its safety envelope with instructors.", priority_num_id)
add_number("Deliver WebAR plus an equivalent fallback using the existing action and grading model.", priority_num_id)
add_number("Make the lesson reliable on low-end devices and weak connectivity before adding more content.", priority_num_id)
add_number("Standardize evidence and run a rules-based adaptive pilot before considering machine learning.", priority_num_id)
add_number("Collect actual usability, acceptance, learning, practical-transfer, and operational evidence before scaling claims or trades.", priority_num_id)

add_heading("10. Pilot evaluation design", 1)
add_para(
    "The dissertation's TAM, SUS, and pre/post design is a useful core, but a technical-school pilot should "
    "also measure physical transfer, safety, fallback parity, device reliability, and instructor workload."
)
add_table(
    ["Evaluation question", "Measure", "Minimum evidence"],
    [
        ("Does it improve learning?", "Aligned pre/post knowledge test, normalized gain, paired analysis, effect size.", "Scores linked to one consented pilot cohort and one stable scenario version."),
        ("Does it improve practical performance?", "Instructor-blinded physical task rubric: sequence, accuracy, safety, time, and help required.", "A workshop transfer assessment independent of the AR app score."),
        ("Is it usable and accepted?", "Standard 10-item SUS; TAM usefulness, ease, enjoyment, and intention; instructor interviews.", "Instrument reliability and item-level reporting, not only averages."),
        ("Does AR work technically?", "Recognition success, false placement, latency, tracking loss, re-scans, crash rate, fallback rate.", "Results by device tier, browser, and network profile."),
        ("Is access equitable?", "Outcome and completion parity across AR and fallback modes and device tiers.", "No grading penalty caused by tracking failure; documented accommodations."),
        ("Can schools operate it?", "Download time, sync delay, support incidents, instructor preparation time, storage use, and recovery success.", "A runbook tested by school staff without developer intervention."),
    ],
    [2250, 4100, 3010],
    font_size=8.7,
)

add_heading("Suggested pilot gates", 2)
add_bullet("Ethics and privacy approval, informed consent, withdrawal process, pseudonymization, retention schedule, and restricted research export are in place.")
add_bullet("Every learner can choose or be assigned a non-AR equivalent without academic penalty.")
add_bullet("No unresolved safety-critical defect; instructors retain authority to stop or invalidate an activity.")
add_bullet("Recognition, fallback, and sync telemetry are operational before recruitment begins.")
add_bullet("Analysis thresholds and exclusion rules are preregistered before outcomes are inspected.")

doc.add_page_break()
add_heading("11. Risks and controls", 1)
add_table(
    ["Risk", "Why it matters", "Control"],
    [
        ("AR novelty mistaken for learning", "High engagement may not transfer to skill.", "Use aligned pre/post tests and an independent physical performance rubric."),
        ("Tracking failure treated as learner error", "Creates unfair scores and damages trust.", "Separate recognition events from competency events; allow fallback and instructor review."),
        ("Unsafe real-equipment use", "Overlays can distract or be spatially wrong.", "Start with training rigs; validate anchors; show stop conditions; require supervision."),
        ("High asset/network cost", "Excludes the exact learners the research intends to serve.", "Budgets, progressive loading, caching, quality tiers, offline completion, and field network tests."),
        ("Premature machine learning", "Small biased datasets can personalize incorrectly and obscure decisions.", "Use versioned transparent rules first; review fairness before any learned model."),
        ("Overclaiming evidence", "The dissertation has no pilot results and OPedu has no AR field trial.", "Use precise product language and publish limitations with every pilot result."),
        ("Fragmented architecture", "New services can increase operational burden for schools.", "Keep the modular monolith; add adapters and background jobs only when evidence requires them."),
    ],
    [1900, 3300, 4160],
    font_size=8.7,
)

add_heading("12. Decisions required before implementation", 1)
add_para("These are product, research, curriculum, and institutional decisions rather than coding tasks:")
add_bullet("Name the first pilot school(s), qualified technical instructor, curriculum standard, and approved training equipment.")
add_bullet("Choose the first supported Android device/browser baseline and whether a marker is acceptable for the first pilot.")
add_bullet("Define whether the pilot requires integration with an existing LMS or LRS; avoid building both without an institutional need.")
add_bullet("Approve the minimum learner data set, lawful basis, consent materials, retention, and research export process.")
add_bullet("Approve success thresholds for SUS, learning gain, practical transfer, recognition reliability, offline synchronization, accessibility, and instructor workload.")
add_bullet("Approve product terminology: '3D technical learning platform' now; 'WebAR platform' only after the qualification gate.")

doc.add_page_break()
add_heading("13. Final assessment", 1)
add_para(
    "OPedu has progressed beyond the dissertation's simplified data and assessment specification in several "
    "important areas: tenant-aware institutional roles, content governance, versioned simulation packages, "
    "deterministic grading, safety evidence, and competency results are already concrete. This is valuable "
    "because these capabilities are difficult to retrofit after an AR demo is built."
)
add_para(
    "The missing work is the research's distinctive contribution: a genuine low-bandwidth WebAR client, "
    "adaptive sequencing, xAPI-compatible analytics, equivalent fallback, and empirical evaluation. The "
    "recommended strategy is therefore not a rewrite. It is a disciplined extension of the current shared "
    "scenario and evidence model into AR, followed by an ethically governed technical-school pilot."
)
add_callout(
    "Decision recommendation",
    "Proceed to an AR minimum vertical slice only after the first procedure, training rig, device baseline, fallback equivalence, data map, and pilot success criteria are approved. Build one evidence-complete battery AR lesson before expanding to additional trades or headsets.",
    PALE_GREEN,
)

doc.add_page_break()
add_heading("Appendix A. Detailed requirement checklist", 1)
add_table(
    ["ID", "Requirement", "Current", "Target acceptance evidence"],
    [
        ("AR-01", "Camera-based browser AR session", "Missing", "Recorded run on each supported browser/device."),
        ("AR-02", "Marker or surface recognition", "Missing", "Accuracy, latency, false-placement, and recovery report."),
        ("AR-03", "Stable spatial placement", "Missing", "Anchor remains usable through the defined task envelope."),
        ("AR-04", "Shared desktop/AR action vocabulary", "Foundation", "Identical action codes and grading outcomes in parity tests."),
        ("AR-05", "Equivalent 2D fallback", "Partial", "All required steps and evidence complete without AR."),
        ("PERF-01", "Lowest-device performance budget", "Missing", "Load, frame-rate, memory, and thermal test results."),
        ("OFF-01", "Offline lesson and evidence", "Missing", "Outage completion and exactly-once synchronization test."),
        ("DATA-01", "xAPI-compatible identifiers", "Planned", "Stable URIs and versioned statement mapping."),
        ("DATA-02", "Idempotent event ingestion", "Missing", "Duplicate and out-of-order delivery tests."),
        ("ADAPT-01", "Explainable sequencing policy", "Missing", "Versioned recommendation and human-readable rationale."),
        ("AUTH-01", "Low/no-code content authoring", "Partial", "Instructor publishes a validated scenario without code."),
        ("EVAL-01", "TAM and SUS instruments", "Missing", "Consent-aware responses and reliability calculation."),
        ("EVAL-02", "Pre/post knowledge evaluation", "Missing", "Aligned instrument and analysis plan."),
        ("EVAL-03", "Physical skill transfer", "Missing", "Independent instructor rubric and scored workshop task."),
        ("SAFE-01", "AR safety review", "Missing", "Instructor approval, stop conditions, incident procedure."),
        ("OPS-01", "School operations readiness", "Missing", "Download, device, charging, storage, support, and recovery runbook."),
    ],
    [750, 2800, 1200, 4610],
    font_size=8.5,
)

add_heading("Appendix B. Sources and repository evidence", 1)
add_heading("Primary research source", 2)
add_para(
    "25093891 AR ELearning Dissertation, 48 pages. Key evidence used: abstract and objectives (pp. 12-15); "
    "conceptual framework and gap analysis (pp. 23-24); methodology, data, adaptive formula, metrics, and "
    "ethics (pp. 25-29); architecture and implementation specification (pp. 30-41); planned results, "
    "limitations, conclusions, recommendations, and future work (pp. 41-44); sample schedule and TAM/SUS "
    "items (p. 48)."
)
add_para(
    "The dissertation's cited studies were not independently re-verified for this repository evaluation. "
    "Several bibliography entries use publication-platform labels rather than complete citations; formal "
    "academic or procurement use should verify each underlying source."
)

add_heading("Repository evidence", 2)
for path in (
    "readme.md",
    "frontend/package.json",
    "frontend/src/pages/SimulationPage.tsx",
    "frontend/src/components/simulation/BatteryWorkshopScene.tsx",
    "mechlab/lab/models.py",
    "mechlab/lab/learning_views.py",
    "docs/architecture/0001-platform-architecture.md",
    "docs/architecture/assessment-evidence.md",
    "docs/architecture/content-versioning.md",
    "docs/product/current-mvp.md",
    "docs/pilot/README.md",
    "EXECUTION_TODO.md",
):
    add_bullet(path)

add_heading("Interpretation note", 2)
add_para(
    "Status labels describe the repository state reviewed on 22 August 2026. 'Implemented' means the core "
    "capability is represented in working application code or authoritative domain models; it does not imply "
    "that a school pilot, accessibility audit, physical-device test, or independent security review has passed."
)

doc.core_properties.title = "OPedu AR Platform Evaluation"
doc.core_properties.subject = "Evaluation against low-bandwidth WebAR e-learning research"
doc.core_properties.author = "OPedu evaluation"
doc.core_properties.keywords = "OPedu, augmented reality, WebAR, technical education, evaluation"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
